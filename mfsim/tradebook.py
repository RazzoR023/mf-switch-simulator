"""Tradebook parsing: Zerodha Coin CSV exports, with a fuzzy fallback for other layouts.

Every transaction comes out as a dict:

    date         ISO date string
    scheme_name  the fund name as the statement spells it
    direction    "buy" or "sell" (anything else is treated as a buy downstream)
    amount       always positive; direction carries the sign
    units, nav, isin, trade_type, txn_type, trade_id   when the source has them
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime

ZERODHA_COLUMNS = {"symbol", "trade_date", "quantity", "price", "trade_type"}

BUY_KEYWORDS = ["purchase", "buy", "sip", "lumpsum", "invest", "switch in"]
SELL_KEYWORDS = ["redeem", "redemption", "sell", "switch out", "withdraw"]

ISO_FORMATS = ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S")
DAYFIRST_FORMATS = (
    "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%d %B %Y",
    "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d-%m-%y", "%d/%m/%y",
)


class TradebookError(Exception):
    """A file that can't be turned into transactions. The message is shown to the user."""


# ─── Cell parsing ────────────────────────────────────────────────────────────

def parse_date(value) -> date | None:
    """ISO first, then day-first (Indian statements); Excel cells arrive as datetimes."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ISO_FORMATS + DAYFIRST_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_number(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if value == value else None
    text = str(value or "").strip().replace(",", "").replace("₹", "").strip()
    if not text:
        return None
    try:
        num = float(text)
    except ValueError:
        return None
    return num if num == num else None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ─── Reading files into header + rows ────────────────────────────────────────

def _read_csv(data: bytes) -> tuple[list[str], list[list]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise TradebookError("The file is empty.")
    return [c.strip() for c in rows[0]], rows[1:]


def _read_xlsx(data: bytes) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise TradebookError(f"Could not read the Excel file: {e}") from e
    try:
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    # Exports often have a title block above the table: the header is the first row
    # mentioning a date and a scheme/NAV/symbol.
    for i, row in enumerate(rows):
        vals = [str(v).lower() for v in row]
        if any("date" in v for v in vals) and any(
            "scheme" in v or "nav" in v or "symbol" in v for v in vals
        ):
            header = [str(c).strip() if c is not None else f"column {j + 1}" for j, c in enumerate(row)]
            return header, rows[i + 1:]
    raise TradebookError(
        "Could not find the header row (a row naming a date column and a scheme/NAV/symbol column)."
    )


def _records(header: list[str], rows: list[list]) -> list[dict]:
    out = []
    for row in rows:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        out.append({h: (row[i] if i < len(row) else None) for i, h in enumerate(header)})
    return out


# ─── Layouts ─────────────────────────────────────────────────────────────────

def _parse_zerodha(records: list[dict]) -> tuple[list[dict], list[str]]:
    """Zerodha Coin tradebook: columns are mapped directly, no guessing."""
    warnings = []
    txns = []
    for r in records:
        d = parse_date(r.get("trade_date"))
        units = parse_number(r.get("quantity"))
        nav = parse_number(r.get("price"))
        units = abs(units) if units is not None else None
        amount = units * nav if units is not None and nav is not None else None
        if d is None or amount is None:
            continue
        txn = {
            "date": d.isoformat(),
            "scheme_name": _text(r.get("symbol")),
            "direction": _text(r.get("trade_type")).lower(),
            "amount": amount,
            "units": units,
            "nav": nav,
            "trade_type": _text(r.get("trade_type")),
        }
        if "isin" in r:
            txn["isin"] = _text(r.get("isin"))
        if "trade_id" in r:
            txn["trade_id"] = _text(r.get("trade_id"))
        txns.append(txn)
    trades = [t for t in txns if t["direction"] in ("buy", "sell")]
    if not trades and txns:
        warnings.append("No buy/sell transactions found, so every row is used.")
        trades = txns
    return trades, warnings


def _parse_fuzzy(records: list[dict], header: list[str]) -> tuple[list[dict], list[str]]:
    """Other exports: guess the columns from their names."""
    col_map: dict[str, str] = {}
    for c in header:
        lc = c.lower()
        if "date" in lc:
            target = "date"
        elif "scheme" in lc or "fund name" in lc:
            target = "scheme_name"
        elif "transaction" in lc or "type" in lc:
            target = "txn_type"
        elif "unit" in lc:
            target = "units"
        elif "nav" in lc or "price" in lc:
            target = "nav"
        elif "amount" in lc or "value" in lc:
            target = "amount"
        else:
            continue
        col_map.setdefault(target, c)

    has_amount = "amount" in col_map or {"units", "nav"} <= col_map.keys()
    missing = [t for t in ("date", "scheme_name") if t not in col_map]
    if not has_amount:
        missing.append("amount")
    if missing:
        raise TradebookError(
            f"Missing columns: {', '.join(missing)}. Found: {', '.join(header)}."
        )

    txns = []
    for r in records:
        d = parse_date(r.get(col_map["date"]))
        units = parse_number(r.get(col_map["units"])) if "units" in col_map else None
        nav = parse_number(r.get(col_map["nav"])) if "nav" in col_map else None
        if "amount" in col_map:
            amount = parse_number(r.get(col_map["amount"]))
        else:
            amount = units * nav if units is not None and nav is not None else None
        if d is None or amount is None:
            continue

        direction = None
        if "txn_type" in col_map:
            txn_type = _text(r.get(col_map["txn_type"])).lower()
            if any(k in txn_type for k in SELL_KEYWORDS):
                direction = "sell"
            elif any(k in txn_type for k in BUY_KEYWORDS):
                direction = "buy"
            elif amount < 0:
                direction = "sell"
            elif amount > 0:
                direction = "buy"
            else:
                continue
        else:
            direction = "buy" if amount >= 0 else "sell"

        txn = {
            "date": d.isoformat(),
            "scheme_name": _text(r.get(col_map["scheme_name"])),
            "direction": direction,
            "amount": abs(amount),
        }
        if "txn_type" in col_map:
            txn["txn_type"] = _text(r.get(col_map["txn_type"]))
        if units is not None:
            txn["units"] = abs(units)
        if nav is not None:
            txn["nav"] = nav
        txns.append(txn)
    return txns, []


def parse_file(name: str, data: bytes) -> tuple[list[dict], list[str]]:
    """Returns (transactions, warnings). Raises TradebookError for an unusable file."""
    lname = name.lower()
    if lname.endswith(".csv"):
        header, rows = _read_csv(data)
    elif lname.endswith(".xlsx"):
        header, rows = _read_xlsx(data)
    else:
        raise TradebookError("Only .csv and .xlsx files are supported.")
    records = _records(header, rows)
    if ZERODHA_COLUMNS <= set(header):
        return _parse_zerodha(records)
    return _parse_fuzzy(records, header)


# ─── Combining files ─────────────────────────────────────────────────────────

def dedupe(txns: list[dict]) -> tuple[list[dict], int, str]:
    """Drop rows that appear in more than one uploaded file. Returns (rows, dropped, method).

    Zerodha rows carry a trade_id, which is used when every row has one. Otherwise a row
    is identified by fund, date and units/NAV (or amount when those are missing).
    """
    columns = set().union(*(t.keys() for t in txns)) if txns else set()
    use_trade_id = bool(txns) and all(
        (t.get("trade_id") or "").strip() not in ("", "nan") for t in txns
    )
    if use_trade_id:
        keycols = ["trade_id"]
        method = "trade ID"
    elif {"units", "nav"} <= columns:
        keycols = [c for c in ("scheme_name", "date", "units", "nav", "trade_type", "direction") if c in columns]
        method = "fund, date, units and NAV"
    else:
        keycols = [c for c in ("scheme_name", "date", "amount", "txn_type", "direction") if c in columns]
        method = "fund, date and amount"

    seen = set()
    out = []
    for t in txns:
        key = tuple(t.get(c) for c in keycols)
        if key in seen:
            continue
        seen.add(key)
        out.append(t)
    return out, len(txns) - len(out), method
