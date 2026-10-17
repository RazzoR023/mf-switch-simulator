"""The Excel report: a Summary sheet and a Transactions sheet, with real numbers (not
pre-formatted text) so the workbook can be sorted, filtered and summed."""

from __future__ import annotations

import io
from datetime import date

MONEY = '"₹"#,##0;-"₹"#,##0'
NAV = '"₹"#,##0.0000'
UNITS = "#,##0.0000"
PCT = '+0.00"%";-0.00"%"'
XIRR = '0.00%'
COUNT = "0"

KIND_FORMATS = {"money": MONEY, "nav": NAV, "units": UNITS, "pct": PCT, "xirr": XIRR, "count": COUNT}


def _notes(row: dict) -> str:
    parts = []
    if row["actual_status"] == "no_nav":
        parts.append("no NAV (actual)")
    if row["alt_status"] == "no_nav":
        parts.append("no NAV (alt)")
    return ", ".join(parts) if parts else "ok"


def build(result: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    green, red = Font(color="15803D", bold=True), Font(color="B91C1C", bold=True)
    header_fill = PatternFill("solid", fgColor="F1F5F9")

    wb = Workbook()

    # ── Summary ──────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    actual_hdr = (
        f"Actual · {result['funds'][0]}" if len(result["funds"]) == 1
        else f"Actual · {len(result['funds'])} funds clubbed"
    )
    alt_hdr = f"Alternate · {result['alternate']['label']}"
    ws.append(["Metric", actual_hdr, alt_hdr])
    for m in result["metrics"]:
        a, b = m["actual"], m["alternate"]
        ws.append([m["label"], a, b])
        r = ws.max_row
        fmt = KIND_FORMATS[m["kind"]]
        for col, val, cls in ((2, a, m["actual_cls"]), (3, b, m["alternate_cls"])):
            cell = ws.cell(row=r, column=col)
            if val is None:
                cell.value = "—"
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.number_format = fmt
            if cls == "better":
                cell.font = green
            elif cls == "worse":
                cell.font = red
    ws.append([])
    ws.append(["Compared on", date.fromisoformat(result["as_of"])])
    ws.cell(row=ws.max_row, column=2).number_format = "dd mmm yyyy"
    if result["exclude_sells"]:
        ws.append(["Sells", "excluded (all bought units treated as still held)"])
    for cell in ws[1]:
        cell.font, cell.fill = bold, header_fill
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 44
    ws.freeze_panes = "A2"

    # ── Transactions ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("Transactions")
    headers = ["Date", "Fund", "Type", "Amount", "Actual NAV", "Actual Units",
               "Alt NAV", "Alt Units", "Notes"]
    ws.append(headers)
    for row in result["detail"]:
        sign = -1 if row["direction"] == "sell" else 1
        ws.append([
            date.fromisoformat(row["date"]), row["fund"], row["direction"].title(),
            sign * row["amount"], row["actual_nav"], row["actual_units"],
            row["alt_nav"], row["alt_units"], _notes(row),
        ])
    formats = {1: "dd mmm yyyy", 4: MONEY, 5: NAV, 6: UNITS, 7: NAV, 8: UNITS}
    for r in range(2, ws.max_row + 1):
        for col, fmt in formats.items():
            ws.cell(row=r, column=col).number_format = fmt
    widths = [13, 48, 7, 14, 14, 14, 14, 14, 26]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill = bold, header_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
